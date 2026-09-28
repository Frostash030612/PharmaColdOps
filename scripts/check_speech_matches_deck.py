"""Do the deck's spoken-worthy claims all appear in the speech script?

For each slide, list the facts the speaker must be able to defend and check that
the speech script states them too. This is the "did I sync the speech?" check.
"""
import sys
from pathlib import Path

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(r"D:\2026第一学期\Group Project")
DECK = ROOT / "proposal" / "PharmaColdOps .pptx"
SPEECH = ROOT / "proposal" / "演讲稿-4人版.md"

deck_slides = []
for slide in Presentation(str(DECK)).slides:
    deck_slides.append("\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame))
speech = SPEECH.read_text(encoding="utf-8")

#: (slide number, label, phrase that must exist in the speech)
checks = [
    (1, "cover: group number", "Project Group 41"),
    (1, "cover: members", "Wang Lepeng"),
    #: the presentation date lives on the cover and is deliberately not read aloud
    (2, "problem: four outcomes", "release, quarantine, retest, scrap"),
    (3, "gap: not first", "we are not first"),
    (4, "solution: reshipment is not a fifth disposition", "not a fifth disposition"),
    (4, "solution: ML stays out of the API", "decision API"),
    (5, "positioning: Controlant", "Controlant"),
    (5, "positioning: foundations", "Solomon"),
    (6, "data: synthetic silent-failure set", "8,000"),
    (6, "data: vaccine distribution", "26,700"),
    (6, "data: thresholds are assumptions", "engineering assumptions"),
    (7, "built: endpoint count", "23 HTTP endpoints"),
    (7, "built: dispatch operations", "15 of them dispatch operations"),
    (7, "built: vue components", "23 components"),
    (7, "built: archived cases and dispatch ops", "25 archived cases and 56 persisted dispatch operations"),
    (7, "built: tests", "305 test functions across 34 test files"),
    (8, "results: hero value", "four dispositions"),
    (8, "results: dispositions named", "release, quarantine, retest or scrap"),
    (8, "results: 54/57 agreement", "54/57"),
    (8, "results: kappa", "0.6434"),
    (8, "results: routing", "133.81 km"),
    (8, "results: routing gain", "6.19%"),
    (8, "results: risk headline", "0.691"),
    (8, "results: majority baseline", "0.161"),
    (8, "results: 19-node network", "19-node Singapore road network"),
    (8, "results: declared limits", "no savings claimed"),
    (9, "limits: quarantine recall", "quarantine"),
    (9, "limits: cause classifier below baseline", "below the baseline"),
    (9, "limits: team roles", "cross-review"),
    (10, "conclusion: four technique groups", "all four IRS technique groups"),
    (10, "conclusion: boundary", "does **not** replace the quality lead's approval"),
]

failures = 0
for slide_no, label, phrase in checks:
    on_slide = phrase.split("**")[-1] in deck_slides[slide_no - 1] or phrase in deck_slides[slide_no - 1]
    in_speech = phrase in speech
    ok = in_speech
    if not ok:
        failures += 1
    print(f"  [{'ok ' if ok else 'BAD'}] slide {slide_no:2d} · {label}"
          f"{'' if ok else '  <- missing from the speech script'}"
          f"{'' if on_slide else '   (note: not literally on the slide either)'}")

print(f"\n{len(checks) - failures}/{len(checks)} deck facts are also stated in the speech")
sys.exit(1 if failures else 0)
