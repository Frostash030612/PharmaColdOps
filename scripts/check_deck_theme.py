"""Verify the added slides now use the original deck's fonts and palette."""
import sys
from collections import Counter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation

prs = Presentation("proposal/PharmaColdOps-Proposal-Presentation-Final.pptx")

THEME_FONTS = {"Libre Baskerville", "DM Sans", "DM Sans Bold"}
THEME_INK = {"454240", "5C4E3D", "B88E23", "063E5F", "F7EDD4", "DDD3BA", "FFFDFA", "FFFFFF"}

print(f"{'页':>3}  {'字体分布':<46} {'字色是否全用主题色':<20}")
for index, slide in enumerate(prs.slides, 1):
    fonts = Counter()
    off_palette = Counter()
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                fonts[run.font.name or "(继承)"] += 1
                try:
                    if run.font.color and run.font.color.type is not None and run.font.color.rgb:
                        rgb = str(run.font.color.rgb)
                        if rgb not in THEME_INK:
                            off_palette[rgb] += 1
                except Exception:
                    pass
    font_note = ", ".join(f"{k}×{v}" for k, v in fonts.most_common(3))
    verdict = "是" if not off_palette else f"否: {dict(off_palette)}"
    print(f"{index:>3}  {font_note:<46} {verdict:<20}")

print("\n=== 新页形状填充色（应为奶油/金/深蓝系）===")
for index, slide in enumerate(prs.slides, 1):
    fills = Counter()
    for shape in slide.shapes:
        try:
            if shape.fill.type == 1 and shape.fill.fore_color and shape.fill.fore_color.rgb:
                fills[str(shape.fill.fore_color.rgb)] += 1
        except Exception:
            pass
    if fills:
        print(f"  第 {index:>2} 页: {dict(fills)}")
