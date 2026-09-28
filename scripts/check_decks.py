"""Validate both decks after the screenshot and group-number edits."""
import sys
import xml.dom.minidom as minidom
import zipfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation

DECKS = [
    Path("proposal/PharmaColdOps .pptx"),
    Path("proposal/PharmaColdOps-Proposal-Presentation-Final.pptx"),
]

for deck in DECKS:
    z = zipfile.ZipFile(deck)
    intact = z.testzip() is None
    parsed = 0
    for name in z.namelist():
        if name.endswith((".xml", ".rels")):
            minidom.parseString(z.read(name))
            parsed += 1
    prs = Presentation(str(deck))
    group = ""
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame and "Project Group" in shape.text_frame.text:
            group = next((l.strip() for l in shape.text_frame.text.splitlines()
                          if "Project Group" in l), "")
            break
    print(f"{deck.name}")
    print(f"   zip={'ok' if intact else 'CORRUPT'} | XML parts parsed={parsed} | "
          f"slides={len(prs.slides)} | {deck.stat().st_size / 1024 / 1024:.2f} MB")
    print(f"   title: {group[:88]}")
