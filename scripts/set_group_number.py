"""Update the project group number in the decks' title slides.

Runs are rewritten in place (keeping their font) rather than through text-frame
.text, which would drop per-run formatting — the title slide's body line mixes a
bold group name with the member list.

    python scripts/set_group_number.py 41 [deck.pptx ...]
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]

NEW = sys.argv[1] if len(sys.argv) > 1 else "41"
DECKS = ([Path(a) for a in sys.argv[2:]] if len(sys.argv) > 2
         else [ROOT / "proposal" / "PharmaColdOps .pptx",
               ROOT / "proposal" / "PharmaColdOps-Proposal-Presentation-Final.pptx"])


def main() -> None:
    for deck in DECKS:
        if not deck.exists():
            print(f"跳过（不存在）: {deck.name}")
            continue
        prs = Presentation(str(deck))
        changed = []
        for index, slide in enumerate(prs.slides, 1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if "Project Group" not in run.text:
                            continue
                        # Keep everything else in the run exactly as it is.
                        import re
                        new_text = re.sub(r"Project Group \d+",
                                          f"Project Group {NEW}", run.text)
                        if new_text != run.text:
                            changed.append((index, shape.name, run.text[:60], new_text[:60]))
                            run.text = new_text
        if changed:
            prs.save(str(deck))
        print(f"\n{deck.name}: {len(changed)} 处改动")
        for idx, name, old, new in changed:
            print(f"   第 {idx} 页 [{name}]")
            print(f"      旧: {old}")
            print(f"      新: {new}")
        if not changed:
            print("   （未发现 Project Group <数字>，无需改动）")


if __name__ == "__main__":
    main()
