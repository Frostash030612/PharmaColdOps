"""Dump PPTX slide text + speaker notes to a text file.

Usage: python dump_pptx.py <deck.pptx> <out.txt>
"""
import sys

from pptx import Presentation
from pptx.util import Emu


def shape_text(shape, indent="  "):
    lines = []
    if shape.shape_type is not None and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            txt = "".join(run.text for run in para.runs)
            if txt.strip():
                lines.append(f"{indent}[lvl{para.level}] {txt}")
    if getattr(shape, "has_table", False) and shape.has_table:
        tbl = shape.table
        lines.append(f"{indent}<TABLE {len(tbl.rows)}x{len(tbl.columns)}>")
        for r in tbl.rows:
            cells = [c.text.replace("\n", " / ").strip() for c in r.cells]
            lines.append(f"{indent}  | " + " | ".join(cells) + " |")
    if shape.shape_type == 6:  # GROUP
        for sub in shape.shapes:
            lines.extend(shape_text(sub, indent + "  "))
    return lines


def main():
    path, out = sys.argv[1], sys.argv[2]
    prs = Presentation(path)
    buf = []
    buf.append(f"### DECK: {path}")
    buf.append(f"slide_size: {prs.slide_width} x {prs.slide_height} "
               f"({Emu(prs.slide_width).inches:.2f}in x {Emu(prs.slide_height).inches:.2f}in)")
    buf.append(f"slide_count: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides, 1):
        buf.append("")
        buf.append(f"===== SLIDE {i} =====")
        try:
            layout = slide.slide_layout.name
        except Exception:
            layout = "?"
        buf.append(f"[layout: {layout}]")
        for shape in slide.shapes:
            kind = str(shape.shape_type)
            if shape.has_text_frame and shape.text_frame.text.strip():
                buf.append(f"  <shape '{shape.name}' {kind}>")
                buf.extend(shape_text(shape, indent="    "))
            elif getattr(shape, "has_table", False) and shape.has_table:
                buf.append(f"  <table shape '{shape.name}'>")
                buf.extend(shape_text(shape, indent="    "))
            elif shape.shape_type == 13:
                buf.append(f"  <PICTURE '{shape.name}'>")
            elif shape.shape_type == 6:
                buf.append(f"  <GROUP '{shape.name}'>")
                buf.extend(shape_text(shape, indent="    "))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            buf.append("  --- NOTES ---")
            for line in slide.notes_slide.notes_text_frame.text.splitlines():
                if line.strip():
                    buf.append(f"    {line}")
    text = "\n".join(buf)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write(text)
    print(f"{path}: slides={len(prs.slides)} chars={len(text)} -> {out}")


if __name__ == "__main__":
    main()
