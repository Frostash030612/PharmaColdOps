"""Compare the ORIGINAL deck's text with my replacement, per text box.

The original string length is the design's own budget for that box (the designer
fitted it), so any box where I wrote more is a candidate for the overflow the
user saw. Sorted by how much longer mine is.
"""
from __future__ import annotations

import difflib
import sys

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SRC = "proposal/PharmaColdOps_tune.pptx"
DST = "proposal/PharmaColdOps-Proposal-Presentation-Final.pptx"


def collect(path):
    out = {}
    prs = Presentation(path)
    for idx, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out[(idx, sh.name)] = sh.text_frame.text
    return out


orig = collect(SRC)
mine = collect(DST)

rows = []
for key, new in mine.items():
    if key not in orig:
        continue                                  # a slide/box I added
    old = orig[key]
    if old.strip() == new.strip():
        continue
    rows.append((len(new) - len(old), len(old), len(new), key, old, new))

rows.sort(key=lambda r: (r[0] / max(r[1], 1)), reverse=True)
print(f"{'Δ%':>6} {'Δchars':>7} {'orig':>5} {'mine':>5}  slide/shape")
for delta, lo, ln, (idx, name), old, new in rows:
    pct = delta / max(lo, 1) * 100
    flag = "  <-- 超出原文预算" if delta > 0 else ""
    print(f"{pct:>+5.0f}% {delta:>+7} {lo:>5} {ln:>5}  {idx:>2}/{name:<10}{flag}")

print("\n" + "=" * 100)
print("逐条对照（只列比原文长的，按溢出风险排序）")
for delta, lo, ln, (idx, name), old, new in [r for r in rows if r[0] > 0]:
    print(f"\n--- slide {idx} / {name}  (+{delta} chars, {lo} -> {ln}) ---")
    print(f"  原文: {old}")
    print(f"  现在: {new}")
